"""Unit tests for generate_docx.py."""

import json
import pytest
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor

from generate_docx import (
    add_bottom_border,
    set_spacing,
    add_run,
    setup_document,
    generate_resume,
    generate_cover_letter,
    FONT_NAME,
    FONT_SIZE_BODY,
    FONT_SIZE_NAME,
    FONT_SIZE_SECTION,
    COLOR_HEADING,
    COLOR_BODY,
    COLOR_META,
    MARGIN_CM,
)


# ─── Helper Functions ─────────────────────────────────────────────────────────


class TestSetupDocument:
    """Tests for the document setup helper."""

    def test_returns_document(self):
        doc = setup_document()
        # Document() is a function returning a Document object; check type by module
        from docx.document import Document as DocumentClass
        assert isinstance(doc, DocumentClass)

    def test_sets_margins(self):
        doc = setup_document()
        section = doc.sections[0]
        assert section.top_margin == Cm(MARGIN_CM)
        assert section.bottom_margin == Cm(MARGIN_CM)
        assert section.left_margin == Cm(MARGIN_CM)
        assert section.right_margin == Cm(MARGIN_CM)

    def test_clears_headers_and_footers(self):
        doc = setup_document()
        section = doc.sections[0]
        # Headers/footers should be empty
        for p in section.header.paragraphs:
            assert p.text == ""
        for p in section.footer.paragraphs:
            assert p.text == ""

    def test_removes_default_paragraph(self):
        doc = setup_document()
        # The default empty paragraph should be removed
        assert len(doc.paragraphs) == 0


class TestAddRun:
    """Tests for the add_run helper."""

    def test_adds_text(self):
        doc = Document()
        p = doc.add_paragraph()
        add_run(p, "Hello World")
        assert p.text == "Hello World"

    def test_sets_font_name(self):
        doc = Document()
        p = doc.add_paragraph()
        run = add_run(p, "Test")
        assert run.font.name == FONT_NAME

    def test_sets_font_size(self):
        doc = Document()
        p = doc.add_paragraph()
        run = add_run(p, "Test", size=12)
        assert run.font.size == Pt(12)

    def test_sets_bold(self):
        doc = Document()
        p = doc.add_paragraph()
        run = add_run(p, "Bold text", bold=True)
        assert run.bold is True

    def test_not_bold_by_default(self):
        doc = Document()
        p = doc.add_paragraph()
        run = add_run(p, "Normal text")
        assert run.bold is False

    def test_sets_italic(self):
        doc = Document()
        p = doc.add_paragraph()
        run = add_run(p, "Italic text", italic=True)
        assert run.italic is True

    def test_sets_color(self):
        doc = Document()
        p = doc.add_paragraph()
        run = add_run(p, "Colored", color=(0xFF, 0x00, 0x00))
        assert run.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)

    def test_uses_default_color(self):
        doc = Document()
        p = doc.add_paragraph()
        run = add_run(p, "Default color")
        assert run.font.color.rgb == RGBColor(*COLOR_BODY)


class TestSetSpacing:
    """Tests for the set_spacing helper."""

    def test_sets_space_before(self):
        doc = Document()
        p = doc.add_paragraph()
        set_spacing(p, before=12)
        assert p.paragraph_format.space_before == Pt(12)

    def test_sets_space_after(self):
        doc = Document()
        p = doc.add_paragraph()
        set_spacing(p, after=8)
        assert p.paragraph_format.space_after == Pt(8)

    def test_defaults_to_zero(self):
        doc = Document()
        p = doc.add_paragraph()
        set_spacing(p)
        assert p.paragraph_format.space_before == Pt(0)
        assert p.paragraph_format.space_after == Pt(0)


class TestAddBottomBorder:
    """Tests for the bottom border helper."""

    def test_adds_border_element(self):
        doc = Document()
        p = doc.add_paragraph("Section Title")
        add_bottom_border(p)
        # The paragraph should have a pBdr element in its properties
        pPr = p._p.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr"
        )
        assert pPr is not None
        pBdr = pPr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr"
        )
        assert pBdr is not None


# ─── Resume Generation ────────────────────────────────────────────────────────


class TestGenerateResume:
    """Tests for full resume document generation."""

    def test_creates_file(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        assert output.exists()
        assert output.stat().st_size > 0

    def test_is_valid_docx(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        assert len(doc.paragraphs) > 0

    def test_contains_name(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "JANE DOE" in all_text  # Name is uppercased

    def test_contains_contact_info(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "jane@example.com" in all_text
        assert "(555) 123-4567" in all_text
        assert "Denver, CO" in all_text

    def test_contains_headline(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Senior Full-Stack Engineer" in all_text

    def test_contains_summary(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "10+ years delivering scalable" in all_text

    def test_contains_skills_section(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Skills" in all_text
        assert "TypeScript" in all_text
        assert "Python" in all_text

    def test_skills_pipe_separated(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        skills_paras = [
            p.text for p in doc.paragraphs
            if "Languages:" in p.text or "Frameworks:" in p.text
        ]
        # Skills should be pipe-separated
        for sp in skills_paras:
            assert " | " in sp

    def test_contains_experience(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Acme Corp" in all_text
        assert "Senior Engineer" in all_text
        assert "Startup Inc" in all_text

    def test_experience_has_date_range(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "08/2022" in all_text
        assert "Present" in all_text

    def test_bullets_have_prefix(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        bullet_paras = [p.text for p in doc.paragraphs if p.text.startswith("•")]
        assert len(bullet_paras) >= 3  # At least 3 bullets total

    def test_contains_education(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Education" in all_text
        assert "University of Colorado" in all_text

    def test_contains_projects(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Projects" in all_text
        assert "OpenSource CLI" in all_text

    def test_contains_certifications(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Certifications" in all_text
        assert "AWS Solutions Architect" in all_text

    def test_creates_parent_dirs(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "nested" / "dir" / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        assert output.exists()

    def test_uses_ats_font(self, tmp_dir, sample_resume_content):
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        # Check first paragraph's first run uses the correct font
        for p in doc.paragraphs:
            if p.runs:
                assert p.runs[0].font.name == FONT_NAME
                break

    def test_name_is_centered(self, tmp_dir, sample_resume_content):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        doc = Document(str(output))
        # First paragraph (name) should be centered
        name_para = doc.paragraphs[0]
        assert name_para.alignment == WD_ALIGN_PARAGRAPH.CENTER

    def test_handles_missing_optional_fields(self, tmp_dir):
        """Resume generation works with only required fields."""
        minimal_content = {
            "personalInfo": {
                "name": "Test User",
                "email": "test@example.com",
            },
            "experience": [
                {
                    "company": "Test Co",
                    "title": "Engineer",
                    "startDate": "01/2020",
                    "endDate": "Present",
                    "bullets": ["Did work"],
                }
            ],
        }
        output = tmp_dir / "minimal_resume.docx"
        generate_resume(minimal_content, Path("/unused"), output)
        assert output.exists()
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "TEST USER" in all_text
        assert "Test Co" in all_text

    def test_handles_empty_skills(self, tmp_dir, sample_resume_content):
        """No crash when skills list is empty."""
        sample_resume_content["skills"] = []
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        assert output.exists()

    def test_handles_empty_experience(self, tmp_dir, sample_resume_content):
        """No crash when experience list is empty."""
        sample_resume_content["experience"] = []
        output = tmp_dir / "resume.docx"
        generate_resume(sample_resume_content, Path("/unused"), output)
        assert output.exists()


# ─── Cover Letter Generation ──────────────────────────────────────────────────


class TestGenerateCoverLetter:
    """Tests for cover letter document generation."""

    def test_creates_file(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        assert output.exists()
        assert output.stat().st_size > 0

    def test_is_valid_docx(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        assert len(doc.paragraphs) > 0

    def test_contains_sender_name(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Jane Doe" in all_text

    def test_contains_contact_info(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "jane@example.com" in all_text
        assert "(555) 123-4567" in all_text

    def test_contains_date(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "August 24, 2026" in all_text

    def test_contains_greeting(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Dear Engineering Hiring Manager," in all_text

    def test_contains_body_paragraphs(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Acme Corp is scaling" in all_text
        assert "1M events daily" in all_text
        assert "cares about craft" in all_text

    def test_contains_closing(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Best regards," in all_text

    def test_wider_margins(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        section = doc.sections[0]
        # Allow small EMU rounding tolerance (±300 EMU)
        assert abs(section.left_margin - Cm(2.5)) < 300
        assert abs(section.right_margin - Cm(2.5)) < 300

    def test_creates_parent_dirs(self, tmp_dir, sample_cover_letter_content):
        output = tmp_dir / "nested" / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        assert output.exists()

    def test_default_greeting_when_missing(self, tmp_dir, sample_cover_letter_content):
        """Falls back to 'Dear Hiring Manager,' if greeting not provided."""
        del sample_cover_letter_content["greeting"]
        output = tmp_dir / "cover_letter.docx"
        generate_cover_letter(sample_cover_letter_content, Path("/unused"), output)
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Dear Hiring Manager," in all_text

    def test_handles_missing_optional_sections(self, tmp_dir):
        """Cover letter works with only required fields."""
        minimal_content = {
            "personalInfo": {
                "name": "Test User",
                "email": "test@example.com",
            },
            "hook": "I'm excited about this role.",
            "value": "I have relevant experience.",
        }
        output = tmp_dir / "minimal_cover.docx"
        generate_cover_letter(minimal_content, Path("/unused"), output)
        assert output.exists()
        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Test User" in all_text
        assert "I'm excited about this role." in all_text
